"""
add_forecast_sheets.py
Append 2027-2029 forecast outputs into the main workbook.
"""

import os
from copy import copy

import pandas as pd
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CLEAN = os.path.join(BASE, "data", "clean")
XLSX = os.path.join(BASE, "NYC_Subway_Safety_Analysis.xlsx")

M1_CSV = os.path.join(CLEAN, "forecast_model1_2027_2029.csv")
M2_CSV = os.path.join(CLEAN, "forecast_model2_2027_2029.csv")

SHEET_M1 = "Forecast M1 2027-2029"
SHEET_M2 = "Forecast M2 2027-2029"
SHEET_SUM = "Forecast Summary 2027-2029"


def write_df_sheet(wb, sheet_name: str, df: pd.DataFrame) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    ws.freeze_panes = "A2"
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    # Basic autosize for readability
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for c in col_cells:
            val = "" if c.value is None else str(c.value)
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(12, max_len + 2), 40)


def write_summary_sheet(wb, m1: pd.DataFrame, m2: pd.DataFrame) -> None:
    if SHEET_SUM in wb.sheetnames:
        del wb[SHEET_SUM]
    ws = wb.create_sheet(SHEET_SUM)

    # Parse dates
    m1c = m1.copy()
    m2c = m2.copy()
    m1c["month"] = pd.to_datetime(m1c["month"])
    m2c["month"] = pd.to_datetime(m2c["month"])

    # Core KPI blocks
    ws["A1"] = "Forecast Summary (2027-2029)"
    ws["A3"] = "Model 1 Rows"
    ws["B3"] = int(len(m1c))
    ws["A4"] = "Model 1 Predicted Breaches"
    ws["B4"] = int(m1c["predicted_sla_breach_label"].sum())
    ws["A5"] = "Model 1 Mean Breach Probability"
    ws["B5"] = float(m1c["predicted_sla_breach_prob"].mean())

    ws["D3"] = "Model 2 Rows"
    ws["E3"] = int(len(m2c))
    ws["D4"] = "Model 2 Predicted Hotspots"
    ws["E4"] = int(m2c["predicted_hotspot_label"].sum())
    ws["D5"] = "Model 2 Mean Hotspot Probability"
    ws["E5"] = float(m2c["predicted_hotspot_prob"].mean())

    # Top-risk tables
    ws["A7"] = "Top 20 Elevator Risks (Model 1)"
    top1 = (
        m1c.groupby("equipment_code", as_index=False)["predicted_sla_breach_prob"]
        .mean()
        .sort_values("predicted_sla_breach_prob", ascending=False)
        .head(20)
        .rename(columns={"predicted_sla_breach_prob": "mean_breach_prob"})
    )
    ws.append([])  # row 8 placeholder to keep section spacing predictable
    start_row_top1 = 8
    ws.cell(row=start_row_top1, column=1, value="equipment_code")
    ws.cell(row=start_row_top1, column=2, value="mean_breach_prob")
    r = start_row_top1 + 1
    for row in top1.itertuples(index=False, name=None):
        ws.cell(row=r, column=1, value=row[0])
        ws.cell(row=r, column=2, value=float(row[1]))
        r += 1

    ws["D7"] = "Top 20 Station Hotspot Risks (Model 2)"
    top2 = (
        m2c.groupby("station_name", as_index=False)["predicted_hotspot_prob"]
        .mean()
        .sort_values("predicted_hotspot_prob", ascending=False)
        .head(20)
        .rename(columns={"predicted_hotspot_prob": "mean_hotspot_prob"})
    )
    start_row_top2 = 8
    ws.cell(row=start_row_top2, column=4, value="station_name")
    ws.cell(row=start_row_top2, column=5, value="mean_hotspot_prob")
    r = start_row_top2 + 1
    for row in top2.itertuples(index=False, name=None):
        ws.cell(row=r, column=4, value=row[0])
        ws.cell(row=r, column=5, value=float(row[1]))
        r += 1

    # Monthly trend summary
    ws["A31"] = "Monthly Predicted Positives (2027-2029)"
    m1_monthly = (
        m1c.assign(month=m1c["month"].dt.to_period("M").astype(str))
        .groupby("month", as_index=False)["predicted_sla_breach_label"]
        .sum()
        .rename(columns={"predicted_sla_breach_label": "m1_predicted_breaches"})
    )
    m2_monthly = (
        m2c.assign(month=m2c["month"].dt.to_period("M").astype(str))
        .groupby("month", as_index=False)["predicted_hotspot_label"]
        .sum()
        .rename(columns={"predicted_hotspot_label": "m2_predicted_hotspots"})
    )
    monthly = m1_monthly.merge(m2_monthly, on="month", how="outer").sort_values("month")

    start_row_month = 32
    ws.cell(row=start_row_month, column=1, value="month")
    ws.cell(row=start_row_month, column=2, value="m1_predicted_breaches")
    ws.cell(row=start_row_month, column=3, value="m2_predicted_hotspots")
    rr = start_row_month + 1
    for row in monthly.itertuples(index=False, name=None):
        ws.cell(row=rr, column=1, value=row[0])
        ws.cell(row=rr, column=2, value=int(row[1]))
        ws.cell(row=rr, column=3, value=int(row[2]))
        rr += 1

    # Formatting
    ws.freeze_panes = "A8"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 24

    title_font = copy(ws["A1"].font)
    title_font.bold = True
    title_font.size = 16
    ws["A1"].font = title_font
    for cell in ["A3", "A4", "A5", "D3", "D4", "D5", "A7", "D7", "A31"]:
        cfont = copy(ws[cell].font)
        cfont.bold = True
        ws[cell].font = cfont

    for row_idx in range(3, 6):
        ws.cell(row=row_idx, column=2).number_format = "0.0000" if row_idx == 5 else "0"
        ws.cell(row=row_idx, column=5).number_format = "0.0000" if row_idx == 5 else "0"

    for row_idx in range(9, 29):
        ws.cell(row=row_idx, column=2).number_format = "0.0000"
        ws.cell(row=row_idx, column=5).number_format = "0.0000"


def main() -> None:
    if not os.path.exists(M1_CSV) or not os.path.exists(M2_CSV):
        raise FileNotFoundError(
            "Forecast CSV output not found. Run analysis/predict_2027_2029.py first."
        )

    m1 = pd.read_csv(M1_CSV)
    m2 = pd.read_csv(M2_CSV)

    wb = load_workbook(XLSX)
    write_df_sheet(wb, SHEET_M1, m1)
    write_df_sheet(wb, SHEET_M2, m2)
    write_summary_sheet(wb, m1, m2)
    wb.save(XLSX)

    print(f"Saved workbook: {XLSX}")
    print(f"Added sheets: {SHEET_M1}, {SHEET_M2}, {SHEET_SUM}")


if __name__ == "__main__":
    main()
